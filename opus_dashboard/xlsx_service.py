from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


ORDER_HEADERS = ("Order Number", "Client", "Minimum Delivery %")
OPENING_HEADERS = (
    "Effective Date",
    "Stock Role",
    "Location",
    "Slab/Bay",
    "Order Number",
    "Opening Tonnes",
)
CONTROL_SHEETS = ("Order Master", "Opening Balances")
INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
MAX_EXCEL_ROWS = 1_048_576
MAX_CELL_CHARACTERS = 32_767
POINT_LEVEL_STORAGE_VALUES = {
    "no loading slab",
    "no offloading slab",
    "point-level / no slab",
}


@dataclass(slots=True)
class WorkbookPreview:
    filename: str
    file_sha256: bytes
    import_id: int | None = None
    orders: list[dict[str, Any]] = field(default_factory=list)
    opening_balances: list[dict[str, Any]] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)

    @property
    def valid(self) -> bool:
        return not self.errors


def create_control_template(path: Path) -> Path:
    workbook = Workbook()
    order_sheet = workbook.active
    order_sheet.title = "Order Master"
    order_sheet.append(ORDER_HEADERS)
    order_sheet.append(("ORDER-EXAMPLE", "Client Name", 99.75))

    opening_sheet = workbook.create_sheet("Opening Balances")
    opening_sheet.append(OPENING_HEADERS)
    opening_sheet.append(
        (date.today(), "Origin", "Loading Point", "Bay 1", "ORDER-EXAMPLE", 0)
    )

    for sheet in (order_sheet, opening_sheet):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1C2545")
        for column in sheet.columns:
            width = max(len(str(cell.value or "")) for cell in column) + 3
            sheet.column_dimensions[column[0].column_letter].width = min(width, 32)

    order_sheet["C2"].number_format = "0.000"
    opening_sheet["A2"].number_format = "yyyy-mm-dd"
    opening_sheet["F2"].number_format = "0.000"
    workbook.save(path)
    return path


def preview_control_workbook(path: Path, filename: str | None = None) -> WorkbookPreview:
    content_hash = hashlib.sha256(path.read_bytes()).digest()
    preview = WorkbookPreview(
        filename=filename or path.name,
        file_sha256=content_hash,
    )
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        preview.errors.append(f"Workbook could not be read: {exc}")
        return preview

    try:
        actual_sheets = tuple(workbook.sheetnames)
        missing = [name for name in CONTROL_SHEETS if name not in actual_sheets]
        unknown = [name for name in actual_sheets if name not in CONTROL_SHEETS]
        if missing:
            preview.errors.append("Missing sheets: " + ", ".join(missing))
        if unknown:
            preview.errors.append("Unknown sheets: " + ", ".join(unknown))
        if missing:
            return preview

        order_rows = _worksheet_rows(workbook["Order Master"], ORDER_HEADERS, preview)
        opening_rows = _worksheet_rows(
            workbook["Opening Balances"],
            OPENING_HEADERS,
            preview,
        )
        _parse_orders(order_rows, preview)
        _parse_openings(opening_rows, preview)
    finally:
        workbook.close()
    return preview


def _worksheet_rows(
    worksheet: Any,
    expected_headers: Sequence[str],
    preview: WorkbookPreview,
) -> list[tuple[int, tuple[Any, ...]]]:
    iterator = worksheet.iter_rows(values_only=True)
    first = next(iterator, ())
    actual_headers = tuple(str(value or "").strip() for value in first)
    while actual_headers and not actual_headers[-1]:
        actual_headers = actual_headers[:-1]
    if actual_headers != tuple(expected_headers):
        preview.errors.append(
            f"{worksheet.title} row 1 headers must be: "
            + ", ".join(expected_headers)
        )
        return []
    rows: list[tuple[int, tuple[Any, ...]]] = []
    for row_number, values in enumerate(iterator, start=2):
        row = tuple(values[: len(expected_headers)])
        if not any(value is not None and str(value).strip() for value in row):
            continue
        rows.append((row_number, row))
    return rows


def _parse_orders(
    rows: list[tuple[int, tuple[Any, ...]]],
    preview: WorkbookPreview,
) -> None:
    seen: set[str] = set()
    for row_number, values in rows:
        order = str(values[0] or "").strip()
        client = str(values[1] or "").strip()
        threshold = _decimal(values[2])
        if not order:
            preview.errors.append(f"Order Master row {row_number}: Order Number is required.")
        if not client:
            preview.errors.append(f"Order Master row {row_number}: Client is required.")
        if threshold is None:
            preview.errors.append(
                f"Order Master row {row_number}: Minimum Delivery % is not a number."
            )
        elif Decimal("0") < threshold <= Decimal("2"):
            threshold *= Decimal("100")
        if threshold is not None and not Decimal("0") <= threshold <= Decimal("200"):
            preview.errors.append(
                f"Order Master row {row_number}: Minimum Delivery % must be 0-200."
            )
        key = order.casefold()
        if key and key in seen:
            preview.errors.append(
                f"Order Master row {row_number}: duplicate Order Number {order!r}."
            )
        seen.add(key)
        if order and client and threshold is not None and Decimal("0") <= threshold <= Decimal("200"):
            preview.orders.append(
                {
                    "row_number": row_number,
                    "order_reference": order,
                    "client_name": client,
                    "minimum_delivery_pct": threshold.quantize(Decimal("0.001")),
                }
            )


def _parse_openings(
    rows: list[tuple[int, tuple[Any, ...]]],
    preview: WorkbookPreview,
) -> None:
    seen: set[tuple[date, str, str, str, str]] = set()
    workbook_orders = {
        str(row["order_reference"]).casefold() for row in preview.orders
    }
    for row_number, values in rows:
        effective_date = _date(values[0])
        raw_role = str(values[1] or "").strip()
        stock_role = {
            "origin": "Origin",
            "destination": "Destination",
        }.get(raw_role.casefold())
        location = str(values[2] or "").strip()
        storage = str(values[3] or "").strip()
        if location and storage.casefold() in POINT_LEVEL_STORAGE_VALUES:
            storage = location
        order = str(values[4] or "").strip()
        tonnes = _decimal(values[5])
        if effective_date is None:
            preview.errors.append(
                f"Opening Balances row {row_number}: Effective Date is invalid."
            )
        if stock_role is None:
            preview.errors.append(
                f"Opening Balances row {row_number}: Stock Role must be "
                "Origin or Destination."
            )
        if not location:
            preview.errors.append(
                f"Opening Balances row {row_number}: Location is required."
            )
        if not storage:
            preview.errors.append(
                f"Opening Balances row {row_number}: Slab/Bay is required."
            )
        if not order:
            preview.errors.append(
                f"Opening Balances row {row_number}: Order Number is required."
            )
        elif order.casefold() not in workbook_orders:
            preview.errors.append(
                f"Opening Balances row {row_number}: Order Number {order!r} "
                "is not present in Order Master."
            )
        if tonnes is None:
            preview.errors.append(
                f"Opening Balances row {row_number}: Opening Tonnes is not a number."
            )
        elif tonnes < 0:
            preview.errors.append(
                f"Opening Balances row {row_number}: Opening Tonnes cannot be negative."
            )
        if effective_date and location and storage and order:
            key = (
                effective_date,
                stock_role or raw_role.casefold(),
                location.casefold(),
                storage.casefold(),
                order.casefold(),
            )
            if key in seen:
                preview.errors.append(
                    f"Opening Balances row {row_number}: duplicate opening-balance key."
                )
            seen.add(key)
        if (
            effective_date
            and stock_role
            and location
            and storage
            and order
            and order.casefold() in workbook_orders
            and tonnes is not None
            and tonnes >= 0
        ):
            preview.opening_balances.append(
                {
                    "row_number": row_number,
                    "effective_date": effective_date,
                    "stock_role": stock_role,
                    "location_name": location,
                    "storage_identifier": storage,
                    "order_reference": order,
                    "opening_tonnes": tonnes.quantize(Decimal("0.001")),
                }
            )


def _decimal(value: Any) -> Decimal | None:
    if value is None or str(value).strip() == "":
        return None
    try:
        return Decimal(str(value).strip().replace(",", "."))
    except InvalidOperation:
        return None


def _date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for formatter in (
        date.fromisoformat,
        lambda candidate: datetime.strptime(candidate, "%d/%m/%Y").date(),
        lambda candidate: datetime.strptime(candidate, "%d-%m-%Y").date(),
    ):
        try:
            return formatter(text)
        except ValueError:
            continue
    return None


class XlsxStreamWriter:
    def __init__(self) -> None:
        self.workbook = Workbook(write_only=True)
        self._used_names: set[str] = set()
        self.row_counts: dict[str, int] = {}

    def add_sheet(
        self,
        title: str,
        headers: Sequence[str],
        rows: Iterable[Mapping[str, Any] | Sequence[Any]],
    ) -> None:
        part = 1
        sheet = self._new_sheet(title, part)
        sheet.append(list(headers))
        rows_in_sheet = 1
        total_rows = 0
        for row in rows:
            if rows_in_sheet >= MAX_EXCEL_ROWS:
                self.row_counts[sheet.title] = rows_in_sheet - 1
                part += 1
                sheet = self._new_sheet(title, part)
                sheet.append(list(headers))
                rows_in_sheet = 1
            values = (
                [row.get(header) for header in headers]
                if isinstance(row, Mapping)
                else list(row)
            )
            sheet.append([_excel_value(value) for value in values])
            rows_in_sheet += 1
            total_rows += 1
        self.row_counts[sheet.title] = rows_in_sheet - 1
        if part > 1:
            self.row_counts[title] = total_rows

    def save(self, path: Path) -> Path:
        if not self.workbook.worksheets:
            self.add_sheet("Summary", ("Message",), (("No rows exported",),))
        self.workbook.save(path)
        return path

    def _new_sheet(self, title: str, part: int) -> Any:
        suffix = f" {part}" if part > 1 else ""
        base = INVALID_SHEET_CHARS.sub("-", title).strip("' ") or "Sheet"
        candidate = (base[: 31 - len(suffix)] + suffix)[:31]
        serial = 2
        while candidate.casefold() in self._used_names:
            serial_suffix = f" ({serial})"
            candidate = base[: 31 - len(serial_suffix)] + serial_suffix
            serial += 1
        self._used_names.add(candidate.casefold())
        return self.workbook.create_sheet(candidate)


def _excel_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (dict, list, tuple, set)):
        value = json.dumps(value, ensure_ascii=False, default=str)
    elif isinstance(value, Decimal):
        return float(value)
    elif isinstance(value, datetime) and value.tzinfo is not None:
        value = value.astimezone().replace(tzinfo=None)
    elif not isinstance(value, (str, int, float, bool, date, datetime)):
        value = str(value)
    if isinstance(value, str) and len(value) > MAX_CELL_CHARACTERS:
        marker = " ... [truncated to Excel cell limit]"
        return value[: MAX_CELL_CHARACTERS - len(marker)] + marker
    return value
