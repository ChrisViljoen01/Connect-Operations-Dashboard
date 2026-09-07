from __future__ import annotations

from datetime import date
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

from opus_dashboard.xlsx_service import (
    XlsxStreamWriter,
    create_control_template,
    preview_control_workbook,
)


class XlsxServiceTests(unittest.TestCase):
    def test_control_template_round_trip(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = create_control_template(Path(directory) / "controls.xlsx")
            preview = preview_control_workbook(path)

        self.assertTrue(preview.valid)
        self.assertEqual(
            preview.orders[0]["order_reference"],
            "ORDER-EXAMPLE",
        )
        self.assertEqual(
            preview.opening_balances[0]["effective_date"],
            date.today(),
        )
        self.assertEqual(preview.opening_balances[0]["stock_role"], "Origin")

    def test_control_preview_reports_all_key_validation_errors(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            workbook = Workbook()
            order_sheet = workbook.active
            order_sheet.title = "Order Master"
            order_sheet.append(("Order Number", "Client", "Minimum Delivery %"))
            order_sheet.append(("ORDER-1", "", "invalid"))
            opening_sheet = workbook.create_sheet("Opening Balances")
            opening_sheet.append(
                (
                    "Effective Date",
                    "Stock Role",
                    "Location",
                    "Slab/Bay",
                    "Order Number",
                    "Opening Tonnes",
                )
            )
            opening_sheet.append(
                ("bad-date", "invalid-role", "", "", "UNKNOWN", -1)
            )
            path = Path(directory) / "invalid.xlsx"
            workbook.save(path)

            preview = preview_control_workbook(path)

        self.assertFalse(preview.valid)
        self.assertGreaterEqual(len(preview.errors), 6)
        self.assertTrue(
            any("Client is required" in error for error in preview.errors)
        )
        self.assertTrue(
            any("cannot be negative" in error for error in preview.errors)
        )
        self.assertTrue(
            any("Stock Role must be Origin or Destination" in error for error in preview.errors)
        )

    def test_point_level_storage_uses_location_as_stable_key(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            workbook = Workbook()
            order_sheet = workbook.active
            order_sheet.title = "Order Master"
            order_sheet.append(("Order Number", "Client", "Minimum Delivery %"))
            order_sheet.append(("ORDER-1", "Client", 99.75))
            opening_sheet = workbook.create_sheet("Opening Balances")
            opening_sheet.append(
                (
                    "Effective Date",
                    "Stock Role",
                    "Location",
                    "Slab/Bay",
                    "Order Number",
                    "Opening Tonnes",
                )
            )
            opening_sheet.append(
                (
                    date.today(),
                    "Destination",
                    "Destination A",
                    "Point-level / no slab",
                    "ORDER-1",
                    100,
                )
            )
            path = Path(directory) / "point-level.xlsx"
            workbook.save(path)

            preview = preview_control_workbook(path)

        self.assertTrue(preview.valid)
        self.assertEqual(
            preview.opening_balances[0]["storage_identifier"],
            "Destination A",
        )

    def test_stream_writer_sanitizes_and_serializes_values(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "export.xlsx"
            writer = XlsxStreamWriter()
            writer.add_sheet(
                "Loading/Exit",
                ("Reference", "Payload"),
                ({"Reference": "ORDBULK-1", "Payload": {"answer": 42}},),
            )
            writer.save(path)

            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                self.assertEqual(workbook.sheetnames, ["Loading-Exit"])
                rows = list(
                    workbook["Loading-Exit"].iter_rows(values_only=True)
                )
                self.assertEqual(rows[1], ("ORDBULK-1", '{"answer": 42}'))
            finally:
                workbook.close()
